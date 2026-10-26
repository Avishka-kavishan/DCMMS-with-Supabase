import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
import re

def set_cell_background(cell, fill_hex):
    tcPr = cell._element.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), fill_hex)
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._element.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def set_table_borders(table, color="D3D3D3", sz="4", val="single"):
    tblPr = table._element.xpath('w:tblPr')
    if tblPr:
        tblBorders = OxmlElement('w:tblBorders')
        for border_name in ['top', 'left', 'bottom', 'right', 'insideH']:
            border = OxmlElement(f'w:{border_name}')
            border.set(qn('w:val'), val)
            border.set(qn('w:sz'), sz)
            border.set(qn('w:space'), '0')
            border.set(qn('w:color'), color)
            tblBorders.append(border)
        border = OxmlElement('w:insideV')
        border.set(qn('w:val'), 'none')
        tblBorders.append(border)
        tblPr[0].append(tblBorders)

def create_styled_document():
    doc = docx.Document()

    # Page Margins
    for section in doc.sections:
        section.top_margin = Inches(0.9)
        section.bottom_margin = Inches(0.9)
        section.left_margin = Inches(0.9)
        section.right_margin = Inches(0.9)

    # Styles
    normal_style = doc.styles['Normal']
    normal_style.font.name = 'Calibri'
    normal_style.font.size = Pt(10.5)
    normal_style.font.color.rgb = RGBColor(33, 37, 41)
    normal_style.paragraph_format.line_spacing = 1.2
    normal_style.paragraph_format.space_after = Pt(5)

    # Document Header Title
    title_p = doc.add_paragraph()
    title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_p.paragraph_format.space_after = Pt(2)
    run_t1 = title_p.add_run("📘 DCMMS - විනය කළමනාකරණ සහ අධීක්ෂණ පද්ධතිය\n")
    run_t1.font.size = Pt(17)
    run_t1.font.bold = True
    run_t1.font.color.rgb = RGBColor(16, 44, 87)

    run_t2 = title_p.add_run("සවිස්තරාත්මක පරිශීලක සහ මෙහෙයුම් අත්පොත | Comprehensive User Manual\n")
    run_t2.font.size = Pt(13)
    run_t2.font.bold = True
    run_t2.font.color.rgb = RGBColor(41, 128, 185)

    sub_p = doc.add_paragraph()
    sub_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub_p.paragraph_format.space_after = Pt(14)
    run_s = sub_p.add_run("ශ්‍රී ලංකා ප්‍රජාතාන්ත්‍රික සමාජවාදී ජනරජයේ අධ්‍යාපන අමාත්‍යාංශය | Ministry of Education, Sri Lanka\nසියලුම අමාත්‍යාංශ නිලධාරීන්, ලේකම්වරුන්, ශාඛා ප්‍රධානීන්, විෂය නිලධාරීන් සහ පද්ධති පරිපාලකයින් සඳහා")
    run_s.font.size = Pt(10)
    run_s.font.italic = True
    run_s.font.color.rgb = RGBColor(108, 117, 125)

    # Line separator
    divider = doc.add_paragraph()
    divider.paragraph_format.space_after = Pt(10)
    pBdr = OxmlElement('w:pBdr')
    bottom = OxmlElement('w:bottom')
    bottom.set(qn('w:val'), 'single')
    bottom.set(qn('w:sz'), '12')
    bottom.set(qn('w:space'), '1')
    bottom.set(qn('w:color'), '102C57')
    pBdr.append(bottom)
    divider._element.get_or_add_pPr().append(pBdr)

    with open("DCMMS_Client_User_Guide.md", "r", encoding="utf-8") as f:
        lines = f.readlines()

    in_table = False
    table_rows = []
    in_code_block = False
    code_block_lines = []

    def flush_table():
        nonlocal in_table, table_rows
        if not table_rows:
            in_table = False
            return
        
        num_cols = max(len(r) for r in table_rows)
        tbl = doc.add_table(rows=len(table_rows), cols=num_cols)
        tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
        set_table_borders(tbl)

        for r_idx, row_data in enumerate(table_rows):
            is_header = (r_idx == 0)
            for c_idx in range(num_cols):
                cell = tbl.cell(r_idx, c_idx)
                cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
                set_cell_margins(cell, top=100, bottom=100, left=130, right=130)
                text = row_data[c_idx] if c_idx < len(row_data) else ""
                
                if is_header:
                    set_cell_background(cell, "102C57")
                    p = cell.paragraphs[0]
                    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
                    p.paragraph_format.space_after = Pt(0)
                    r = p.add_run(text.strip())
                    r.font.bold = True
                    r.font.size = Pt(9.5)
                    r.font.color.rgb = RGBColor(255, 255, 255)
                else:
                    if r_idx % 2 == 1:
                        set_cell_background(cell, "F8F9FA")
                    else:
                        set_cell_background(cell, "FFFFFF")
                    p = cell.paragraphs[0]
                    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
                    p.paragraph_format.space_after = Pt(0)
                    parts = re.split(r'(\*\*.*?\*\*)', text.strip())
                    for part in parts:
                        if part.startswith('**') and part.endswith('**'):
                            run = p.add_run(part[2:-2])
                            run.font.bold = True
                            run.font.size = Pt(9.5)
                        elif part.startswith('*') and part.endswith('*'):
                            run = p.add_run(part[1:-1])
                            run.font.italic = True
                            run.font.size = Pt(9.5)
                        else:
                            run = p.add_run(part)
                            run.font.size = Pt(9.5)

        p_after = doc.add_paragraph()
        p_after.paragraph_format.space_after = Pt(6)
        table_rows = []
        in_table = False

    def flush_code_block():
        nonlocal in_code_block, code_block_lines
        if not code_block_lines:
            in_code_block = False
            return
        
        tbl = doc.add_table(rows=1, cols=1)
        tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
        cell = tbl.cell(0, 0)
        set_cell_background(cell, "F1F5F9")
        set_cell_margins(cell, top=120, bottom=120, left=150, right=150)
        
        p = cell.paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.line_spacing = 1.15
        
        full_text = "\n".join(code_block_lines)
        r = p.add_run(full_text)
        r.font.name = "Consolas"
        r.font.size = Pt(9)
        r.font.color.rgb = RGBColor(15, 23, 42)
        
        p_after = doc.add_paragraph()
        p_after.paragraph_format.space_after = Pt(6)
        code_block_lines = []
        in_code_block = False

    skip_header_lines = 4
    current_line = 0

    for line_raw in lines:
        current_line += 1
        line = line_raw.strip()

        if current_line <= skip_header_lines:
            continue

        # Check for code block fences
        if line.startswith('```'):
            if in_code_block:
                flush_code_block()
            else:
                if in_table:
                    flush_table()
                in_code_block = True
                code_block_lines = []
            continue

        if in_code_block:
            code_block_lines.append(line_raw.rstrip())
            continue

        # Check for table row
        if line.startswith('|') and line.endswith('|'):
            if re.match(r'^\|(\s*:?-+:?\s*\|)+$', line):
                continue
            cells = [c.strip() for c in line.split('|')[1:-1]]
            table_rows.append(cells)
            in_table = True
            continue
        elif in_table:
            flush_table()

        if not line:
            continue

        # Heading 2 (e.g. ## 1. පද්ධති හැඳින්වීම)
        if line.startswith('## '):
            h_text = line[3:].strip()
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(16)
            p.paragraph_format.space_after = Pt(5)
            p.paragraph_format.keep_with_next = True
            r = p.add_run(h_text)
            r.font.size = Pt(13)
            r.font.bold = True
            r.font.color.rgb = RGBColor(16, 44, 87)
            continue

        # Heading 3 (e.g. ### පෝරමයේ ක්ෂේත්‍ර)
        if line.startswith('### '):
            h_text = line[4:].strip()
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(12)
            p.paragraph_format.space_after = Pt(4)
            p.paragraph_format.keep_with_next = True
            r = p.add_run(h_text)
            r.font.size = Pt(11.5)
            r.font.bold = True
            r.font.color.rgb = RGBColor(31, 110, 165)
            continue

        # Heading 4 (e.g. #### A. ලැබීමේ ක්‍රමය)
        if line.startswith('#### '):
            h_text = line[5:].strip()
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(9)
            p.paragraph_format.space_after = Pt(3)
            p.paragraph_format.keep_with_next = True
            r = p.add_run(h_text)
            r.font.size = Pt(10.5)
            r.font.bold = True
            r.font.color.rgb = RGBColor(44, 62, 80)
            continue

        # Callout blockquote (e.g. > ⚠️ **අවවාදයයි:** ...)
        if line.startswith('> '):
            box_text = line[2:].strip()
            tbl_box = doc.add_table(rows=1, cols=1)
            tbl_box.alignment = WD_TABLE_ALIGNMENT.CENTER
            cell = tbl_box.cell(0, 0)
            set_cell_background(cell, "FEF9E7" if "⚠️" in box_text else "EBF5FB")
            set_cell_margins(cell, top=120, bottom=120, left=160, right=160)
            
            tcPr = cell._element.get_or_add_tcPr()
            tcBorders = OxmlElement('w:tcBorders')
            leftBdr = OxmlElement('w:left')
            leftBdr.set(qn('w:val'), 'single')
            leftBdr.set(qn('w:sz'), '24')
            leftBdr.set(qn('w:space'), '0')
            leftBdr.set(qn('w:color'), 'F39C12' if "⚠️" in box_text else '2980B9')
            tcBorders.append(leftBdr)
            for side in ['top', 'bottom', 'right']:
                b = OxmlElement(f'w:{side}')
                b.set(qn('w:val'), 'none')
                tcBorders.append(b)
            tcPr.append(tcBorders)

            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            parts = re.split(r'(\*\*.*?\*\*)', box_text)
            for part in parts:
                if part.startswith('**') and part.endswith('**'):
                    run = p.add_run(part[2:-2])
                    run.font.bold = True
                    run.font.size = Pt(10)
                else:
                    run = p.add_run(part)
                    run.font.size = Pt(10)

            p_after = doc.add_paragraph()
            p_after.paragraph_format.space_after = Pt(6)
            continue

        # Checklist item (e.g. - [ ] ...)
        if line.startswith('- [ ] ') or line.startswith('- [x] '):
            is_checked = line.startswith('- [x] ')
            item_text = line[6:].strip()
            p = doc.add_paragraph()
            p.paragraph_format.space_after = Pt(3)
            p.paragraph_format.left_indent = Inches(0.25)
            r_box = p.add_run("☑ " if is_checked else "☐ ")
            r_box.font.bold = True
            r_box.font.size = Pt(11)
            parts = re.split(r'(\*\*.*?\*\*)', item_text)
            for part in parts:
                if part.startswith('**') and part.endswith('**'):
                    run = p.add_run(part[2:-2])
                    run.font.bold = True
                elif part.startswith('*') and part.endswith('*'):
                    run = p.add_run(part[1:-1])
                    run.font.italic = True
                else:
                    run = p.add_run(part)
            continue

        # Bullet list item (- or *)
        if line.startswith('- ') or line.startswith('* '):
            bullet_text = line[2:].strip()
            p = doc.add_paragraph(style='List Bullet')
            p.paragraph_format.space_after = Pt(3)
            p.paragraph_format.left_indent = Inches(0.25)
            parts = re.split(r'(\*\*.*?\*\*)', bullet_text)
            for part in parts:
                if part.startswith('**') and part.endswith('**'):
                    run = p.add_run(part[2:-2])
                    run.font.bold = True
                elif part.startswith('*') and part.endswith('*'):
                    run = p.add_run(part[1:-1])
                    run.font.italic = True
                else:
                    run = p.add_run(part)
            continue

        # Numbered list item (1. 2. etc)
        m_num = re.match(r'^(\d+)\.\s+(.*)$', line)
        if m_num:
            num_str = m_num.group(1)
            num_text = m_num.group(2).strip()
            p = doc.add_paragraph()
            p.paragraph_format.space_after = Pt(3)
            p.paragraph_format.left_indent = Inches(0.25)
            r_num = p.add_run(f"{num_str}. ")
            r_num.font.bold = True
            parts = re.split(r'(\*\*.*?\*\*)', num_text)
            for part in parts:
                if part.startswith('**') and part.endswith('**'):
                    run = p.add_run(part[2:-2])
                    run.font.bold = True
                elif part.startswith('*') and part.endswith('*'):
                    run = p.add_run(part[1:-1])
                    run.font.italic = True
                else:
                    run = p.add_run(part)
            continue

        # Normal paragraph
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(5)
        parts = re.split(r'(\*\*.*?\*\*)', line)
        for part in parts:
            if part.startswith('**') and part.endswith('**'):
                run = p.add_run(part[2:-2])
                run.font.bold = True
            elif part.startswith('*') and part.endswith('*'):
                run = p.add_run(part[1:-1])
                run.font.italic = True
            else:
                run = p.add_run(part)

    if in_table:
        flush_table()
    if in_code_block:
        flush_code_block()

    primary_path = "DCMMS_Full_User_Guide_Handover.docx"
    doc.save(primary_path)
    print(f"Successfully generated {primary_path}")

    fallback_path = "DCMMS_Client_User_Guide.docx"
    try:
        doc.save(fallback_path)
        print(f"Successfully updated {fallback_path}")
    except PermissionError:
        print(f"Note: {fallback_path} is currently open in Microsoft Word. Saved to {primary_path} instead.")

if __name__ == "__main__":
    create_styled_document()
