import os
import re
import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import qn, nsdecls

def set_cell_background(cell, fill_hex):
    tcPr = cell._element.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=120, bottom=120, left=150, right=150):
    tcPr = cell._element.get_or_add_tcPr()
    tcMar = parse_xml(f'''
        <w:tcMar {nsdecls("w")}>
            <w:top w:w="{top}" w:type="dxa"/>
            <w:bottom w:w="{bottom}" w:type="dxa"/>
            <w:left w:w="{left}" w:type="dxa"/>
            <w:right w:w="{right}" w:type="dxa"/>
        </w:tcMar>
    ''')
    tcPr.append(tcMar)

def set_table_borders(table, color="D3D3D3"):
    tblPr = table._element.xpath('w:tblPr')
    if tblPr:
        borders = parse_xml(f'''
            <w:tblBorders {nsdecls("w")}>
                <w:top w:val="single" w:sz="6" w:space="0" w:color="{color}"/>
                <w:bottom w:val="single" w:sz="6" w:space="0" w:color="{color}"/>
                <w:insideH w:val="single" w:sz="4" w:space="0" w:color="{color}"/>
                <w:insideV w:val="single" w:sz="4" w:space="0" w:color="{color}"/>
                <w:left w:val="none"/>
                <w:right w:val="none"/>
            </w:tblBorders>
        ''')
        tblPr[0].append(borders)

def add_callout(doc, text, box_type="NOTE"):
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.autofit = False
    
    color_map = {
        "NOTE": ("EBF8FF", "2B6CB0"),
        "IMPORTANT": ("E6FFFA", "234E52"),
        "WARNING": ("FFFAF0", "C05621"),
        "CAUTION": ("FFF5F5", "C53030")
    }
    bg_hex, border_hex = color_map.get(box_type, ("F7FAFC", "4A5568"))
    
    cell = tbl.cell(0, 0)
    cell.width = Inches(6.5)
    set_cell_background(cell, bg_hex)
    set_cell_margins(cell, top=140, bottom=140, left=180, right=180)
    
    tcPr = cell._element.get_or_add_tcPr()
    borders = parse_xml(f'''
        <w:tcBorders {nsdecls("w")}>
            <w:left w:val="single" w:sz="24" w:space="0" w:color="{border_hex}"/>
            <w:top w:val="none"/>
            <w:right w:val="none"/>
            <w:bottom w:val="none"/>
        </w:tcBorders>
    ''')
    tcPr.append(borders)
    
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = 1.15
    run = p.add_run(text)
    run.font.name = "Arial"
    run.font.size = Pt(10)
    doc.add_paragraph().paragraph_format.space_after = Pt(4)

def build_docx_from_markdown(md_path, docx_path):
    doc = Document()
    
    # Page setup: Standard Letter, 0.8 inch margins
    sections = doc.sections
    for section in sections:
        section.top_margin = Inches(0.8)
        section.bottom_margin = Inches(0.8)
        section.left_margin = Inches(0.8)
        section.right_margin = Inches(0.8)
        
    with open(md_path, "r", encoding="utf-8") as f:
        lines = f.readlines()
        
    i = 0
    total_lines = len(lines)
    
    in_code_block = False
    code_block_lines = []
    
    while i < total_lines:
        line = lines[i]
        stripped = line.strip()
        
        # Check code block
        if stripped.startswith("```"):
            if in_code_block:
                # End of code block
                in_code_block = False
                code_text = "\n".join(code_block_lines)
                add_callout(doc, code_text, box_type="NOTE")
                code_block_lines = []
            else:
                in_code_block = True
                code_block_lines = []
            i += 1
            continue
            
        if in_code_block:
            code_block_lines.append(line.rstrip("\r\n"))
            i += 1
            continue
            
        # Empty line
        if not stripped:
            i += 1
            continue
            
        # Horizontal rule
        if stripped in ["---", "***", "___"]:
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(6)
            p.paragraph_format.space_after = Pt(6)
            i += 1
            continue
            
        # Alert / Blockquote
        if stripped.startswith(">"):
            quote_text = stripped.lstrip("> ").strip()
            box_type = "NOTE"
            if "අවවාදයයි" in quote_text or "WARNING" in quote_text.upper():
                box_type = "WARNING"
            elif "සටහන" in quote_text or "NOTE" in quote_text.upper():
                box_type = "IMPORTANT"
            add_callout(doc, quote_text, box_type=box_type)
            i += 1
            continue
            
        # Table detection
        if "|" in stripped and i + 1 < total_lines and "|" in lines[i+1] and ("---" in lines[i+1] or ":---" in lines[i+1]):
            # Parse table
            table_rows = []
            header_line = stripped
            sep_line = lines[i+1].strip()
            i += 2
            
            headers = [c.strip() for c in header_line.split("|")[1:-1]]
            
            while i < total_lines and "|" in lines[i].strip():
                row_cells = [c.strip() for c in lines[i].strip().split("|")[1:-1]]
                if row_cells:
                    table_rows.append(row_cells)
                i += 1
                
            # Create Docx Table
            num_cols = len(headers)
            tbl = doc.add_table(rows=len(table_rows) + 1, cols=num_cols)
            tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
            set_table_borders(tbl)
            
            # Header Row
            hdr_cells = tbl.rows[0].cells
            for c_idx, title in enumerate(headers):
                hdr_cells[c_idx].text = title
                set_cell_background(hdr_cells[c_idx], "1E3A8A") # Navy
                set_cell_margins(hdr_cells[c_idx], top=120, bottom=120, left=140, right=140)
                p = hdr_cells[c_idx].paragraphs[0]
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT
                for r in p.runs:
                    r.font.name = "Arial"
                    r.font.bold = True
                    r.font.size = Pt(9.5)
                    r.font.color.rgb = RGBColor(255, 255, 255)
                    
            # Data Rows
            for r_idx, row_data in enumerate(table_rows):
                row_cells = tbl.rows[r_idx + 1].cells
                fill_color = "F8FAFC" if r_idx % 2 == 1 else "FFFFFF"
                for c_idx in range(min(num_cols, len(row_data))):
                    cell = row_cells[c_idx]
                    cell_text = row_data[c_idx].replace("<br>", "\n").replace("**", "")
                    cell.text = cell_text
                    set_cell_background(cell, fill_color)
                    set_cell_margins(cell, top=100, bottom=100, left=140, right=140)
                    p = cell.paragraphs[0]
                    p.paragraph_format.line_spacing = 1.15
                    for r in p.runs:
                        r.font.name = "Arial"
                        r.font.size = Pt(9)
                        r.font.color.rgb = RGBColor(30, 41, 59)
                        
            doc.add_paragraph().paragraph_format.space_after = Pt(6)
            continue
            
        # Headings
        if stripped.startswith("# "):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(14)
            p.paragraph_format.space_after = Pt(4)
            run = p.add_run(stripped[2:].replace("**", ""))
            run.font.name = "Arial"
            run.font.size = Pt(20)
            run.font.bold = True
            run.font.color.rgb = RGBColor(30, 58, 138) # Dark Navy
            i += 1
            continue
            
        if stripped.startswith("## "):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(12)
            p.paragraph_format.space_after = Pt(3)
            run = p.add_run(stripped[3:].replace("**", ""))
            run.font.name = "Arial"
            run.font.size = Pt(14)
            run.font.bold = True
            run.font.color.rgb = RGBColor(37, 99, 235) # Royal Blue
            i += 1
            continue
            
        if stripped.startswith("### "):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(8)
            p.paragraph_format.space_after = Pt(2)
            run = p.add_run(stripped[4:].replace("**", ""))
            run.font.name = "Arial"
            run.font.size = Pt(12)
            run.font.bold = True
            run.font.color.rgb = RGBColor(15, 23, 42) # Slate
            i += 1
            continue

        if stripped.startswith("#### "):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(6)
            p.paragraph_format.space_after = Pt(2)
            run = p.add_run(stripped[5:].replace("**", ""))
            run.font.name = "Arial"
            run.font.size = Pt(11)
            run.font.bold = True
            run.font.color.rgb = RGBColor(71, 85, 105) # Slate Medium
            i += 1
            continue

        # Bullet list
        if stripped.startswith("- ") or stripped.startswith("* "):
            p = doc.add_paragraph(style="List Bullet")
            p.paragraph_format.space_before = Pt(1)
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.line_spacing = 1.15
            content = stripped[2:]
            
            # Basic bold parsing
            parts = re.split(r"(\*\*.*?\*\*)", content)
            for part in parts:
                if part.startswith("**") and part.endswith("**"):
                    r = p.add_run(part[2:-2])
                    r.font.name = "Arial"
                    r.font.bold = True
                    r.font.size = Pt(10)
                else:
                    r = p.add_run(part)
                    r.font.name = "Arial"
                    r.font.size = Pt(10)
            i += 1
            continue

        # Numbered list
        match_num = re.match(r"^(\d+)\.\s+(.*)", stripped)
        if match_num:
            p = doc.add_paragraph(style="List Number")
            p.paragraph_format.space_before = Pt(1)
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.line_spacing = 1.15
            content = match_num.group(2)
            
            parts = re.split(r"(\*\*.*?\*\*)", content)
            for part in parts:
                if part.startswith("**") and part.endswith("**"):
                    r = p.add_run(part[2:-2])
                    r.font.name = "Arial"
                    r.font.bold = True
                    r.font.size = Pt(10)
                else:
                    r = p.add_run(part)
                    r.font.name = "Arial"
                    r.font.size = Pt(10)
            i += 1
            continue

        # Checklist list: - [ ]
        if stripped.startswith("- [ ]") or stripped.startswith("- [x]"):
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.25)
            p.paragraph_format.space_before = Pt(1)
            p.paragraph_format.space_after = Pt(2)
            content = "☑ " if "- [x]" in stripped else "☐ "
            content += stripped[5:].strip()
            parts = re.split(r"(\*\*.*?\*\*)", content)
            for part in parts:
                if part.startswith("**") and part.endswith("**"):
                    r = p.add_run(part[2:-2])
                    r.font.name = "Arial"
                    r.font.bold = True
                    r.font.size = Pt(10)
                else:
                    r = p.add_run(part)
                    r.font.name = "Arial"
                    r.font.size = Pt(10)
            i += 1
            continue

        # Normal Paragraph
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(2)
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.line_spacing = 1.15
        parts = re.split(r"(\*\*.*?\*\*)", stripped)
        for part in parts:
            if part.startswith("**") and part.endswith("**"):
                r = p.add_run(part[2:-2])
                r.font.name = "Arial"
                r.font.bold = True
                r.font.size = Pt(10)
            else:
                r = p.add_run(part)
                r.font.name = "Arial"
                r.font.size = Pt(10)
                
        i += 1

    doc.save(docx_path)
    print(f"Successfully generated DOCX at: {docx_path}")

if __name__ == "__main__":
    md_file = r"e:\Avishka Kavishan's Projects\DCMMS-with-Supabase\DCMMS_Client_User_Guide.md"
    docx_file = r"e:\Avishka Kavishan's Projects\DCMMS-with-Supabase\DCMMS_Complete_User_Guide.docx"
    build_docx_from_markdown(md_file, docx_file)
